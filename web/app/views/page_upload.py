"""Upload page: AI smart load, batch upload, directory scan, smart exploration."""

from __future__ import annotations

import json
import sys
import tempfile
from pathlib import Path

_project_root = str(Path(__file__).parent.parent.parent.parent)
if _project_root not in sys.path:
    sys.path.insert(0, _project_root)

import streamlit as st
import pandas as pd

from lmbagent.data.loader import load_auto, detect_format
from lmbagent.data.transformer import add_cycle_summary, split_by_cycles
from lmbagent.data.store import DataStore
from lmbagent.data.catalog import scan_directory, batch_import

_EXPERIENCE_DIR = Path.home() / ".lmbagent" / "load_experiences"


def render_upload_page():
    try:
        _render()
    except Exception as e:
        import streamlit as st
        st.error(f"页面渲染出错: {e}")


def _render():
    store = DataStore()

    tab_llm, tab_multi, tab_scan, tab_explore = st.tabs([
        "AI 智能加载", "批量上传", "目录扫描", "智能探索",
    ])

    with tab_llm:
        _render_ai_load(store)

    with tab_multi:
        _render_batch_upload(store)

    with tab_scan:
        _render_directory_scan(store)

    with tab_explore:
        _render_smart_explore(store)


def _render_ai_load(store: DataStore):
    st.subheader("AI 智能加载与校对")
    st.markdown("""
    上传数据文件，AI 自动分析文件结构并加载。加载后可查看数据预览，
    如有问题可直接用自然语言对话校正，每次校正后预览自动更新。

    **校正示例:** "第一列是循环号，第三列是电压" / "跳过前5行" / "容量需按电流正负拆分"
    """)

    llm_file = st.file_uploader(
        "上传数据文件（< 20 MB）",
        type=["csv", "txt", "tsv", "xlsx", "xls", "npy"],
        key="llm_upload",
    )

    st.caption("大文件（> 20 MB）请使用下方路径加载")
    llm_server_path = st.text_input(
        "服务器文件路径",
        placeholder="/path/to/your/data.csv",
        key="llm_server_path",
    )

    llm_tmp_path = None
    llm_file_name = None

    if llm_server_path and Path(llm_server_path).exists():
        llm_tmp_path = Path(llm_server_path)
        llm_file_name = llm_tmp_path.name
        llm_file_size = llm_tmp_path.stat().st_size
        st.write(f"**文件:** {llm_file_name} ({llm_file_size / 1024:.1f} KB)")
    elif llm_file:
        llm_tmp_dir = Path(tempfile.gettempdir()) / "lmbagent_llm"
        llm_tmp_dir.mkdir(parents=True, exist_ok=True)
        llm_tmp_path = llm_tmp_dir / llm_file.name
        llm_tmp_path.write_bytes(llm_file.getvalue())
        llm_file_name = llm_file.name
        st.write(f"**文件:** {llm_file.name} ({llm_file.size / 1024:.1f} KB)")

    if not (llm_tmp_path and llm_file_name):
        return

    llm_data_id = st.text_input("数据集 ID (可选)", value="",
                                placeholder=llm_file_name.split(".")[0],
                                key="llm_data_id")

    if "llm_messages" not in st.session_state:
        st.session_state.llm_messages = []
    if "llm_analysis" not in st.session_state:
        st.session_state.llm_analysis = None
    if "llm_dataset" not in st.session_state:
        st.session_state.llm_dataset = None
    if "llm_file_key" not in st.session_state:
        st.session_state.llm_file_key = None

    if st.session_state.llm_file_key != llm_file_name:
        st.session_state.llm_messages = []
        st.session_state.llm_analysis = None
        st.session_state.llm_dataset = None
        st.session_state.llm_file_key = llm_file_name

    # --- Load experience selector ---
    saved_experiences = _list_experiences()
    if saved_experiences:
        exp_names = ["(无)"] + [e["name"] for e in saved_experiences]
        selected_exp = st.selectbox("加载上次校对经验", exp_names, key="llm_exp_select")
        if selected_exp != "(无)" and st.button("应用经验", key="llm_apply_exp"):
            for e in saved_experiences:
                if e["name"] == selected_exp:
                    st.session_state.llm_analysis = e["analysis"]
                    _apply_experience_and_load(llm_tmp_path, llm_data_id, e["analysis"])
                    break

    # --- AI analyze button ---
    if st.button("AI 分析并加载", type="primary", key="llm_load_btn"):
        import time as _time
        from lmbagent.data.llm_loader import llm_analyze_file, llm_smart_load

        _t0 = _time.time()
        progress = st.progress(0, text="AI 正在分析文件结构...")
        try:
            analysis = llm_analyze_file(str(llm_tmp_path))
            _t1 = _time.time()
            st.info(f"[计时] 分析: {_t1-_t0:.2f}s  format={analysis.get('format')}  llm={analysis.get('llm_used')}")
        except Exception as e:
            progress.empty()
            st.error(f"文件分析失败: {e}")
            analysis = {"format": "unknown", "reasoning": str(e)}

        if not analysis or analysis.get("format") == "unknown":
            progress.empty()
            st.error(f"AI 无法识别文件格式: {analysis.get('reasoning', '未知错误')}")
            st.session_state.llm_analysis = analysis
            st.session_state.llm_messages.append({
                "role": "assistant",
                "content": f"AI 无法识别文件格式: {analysis.get('reasoning', '未知错误')}\n\n请在下方描述文件格式，我会帮你校正加载。",
            })
        else:
            progress.progress(60, text="格式已识别，正在加载数据...")
            st.session_state.llm_analysis = analysis
            _load_and_update_preview(llm_tmp_path, llm_data_id, analysis, progress)

        st.rerun()

    # --- Raw file preview ---
    with st.expander("原始文件预览", expanded=False):
        if st.button("加载预览", key="load_preview_btn"):
            _show_file_preview(llm_tmp_path)
        else:
            st.info("点击上方按钮加载文件预览")

    # --- Data preview (always visible when dataset exists) ---
    ds_preview = st.session_state.llm_dataset
    if ds_preview is not None:
        _render_data_preview(ds_preview)

    # --- Chat messages ---
    for msg in st.session_state.llm_messages:
        with st.chat_message(msg["role"]):
            st.markdown(msg["content"])

    # --- Multi-turn correction chat ---
    if user_input := st.chat_input("输入校正指令，如: 第一列是循环号，容量需按电流正负拆分..."):
        st.session_state.llm_messages.append({"role": "user", "content": user_input})

        with st.spinner("AI 正在根据你的反馈重新加载..."):
            from lmbagent.data.llm_loader import corrective_load

            ds, analysis = corrective_load(
                str(llm_tmp_path),
                user_feedback=user_input,
                previous_analysis=st.session_state.llm_analysis,
                data_id=llm_data_id or None,
            )

        st.session_state.llm_analysis = analysis

        if ds is not None and ds.num_data_points > 0:
            st.session_state.llm_dataset = ds
            col_info = ", ".join(ds.raw_data.columns.tolist())
            cc_col = ds.raw_data.get("charge_capacity")
            dc_col = ds.raw_data.get("discharge_capacity")
            cc_sum = cc_col.sum() if cc_col is not None else 0
            dc_sum = dc_col.sum() if dc_col is not None else 0
            reply = (
                f"校正加载成功!\n\n"
                f"- 格式: `{analysis.get('format', 'unknown')}`\n"
                f"- 数据: {ds.num_data_points} 行, {ds.num_cycles} 循环\n"
                f"- 列: {col_info}\n"
                f"- 充电容量: {cc_sum:.4f} Ah | 放电容量: {dc_sum:.4f} Ah\n"
            )
            if analysis.get("reasoning"):
                reply += f"- AI说明: {analysis['reasoning']}\n"
            reply += "\n请检查上方数据预览，如需进一步调整请继续输入。"
            st.session_state.llm_messages.append({"role": "assistant", "content": reply})
        else:
            reason = analysis.get("reasoning", "未知原因")
            st.session_state.llm_messages.append({
                "role": "assistant",
                "content": (
                    f"校正后仍未成功加载: {reason}\n\n"
                    "请提供更多格式信息，例如:\n"
                    "- 列名含义\n- 分隔符类型\n- 需要跳过的行数"
                ),
            })

        st.rerun()

    # --- Save to DB + Save experience ---
    if st.session_state.llm_dataset is not None:
        st.divider()
        c1, c2 = st.columns(2)
        with c1:
            if st.button("确认入库", type="primary", key="llm_save_btn"):
                ds = st.session_state.llm_dataset
                ds = add_cycle_summary(ds)
                store.put(ds)
                st.session_state.active_dataset_id = ds.data_id
                st.success(
                    f"已入库! ID: `{ds.data_id}`, "
                    f"{ds.num_data_points} points, {ds.num_cycles} cycles"
                )
                st.session_state.llm_messages = []
                st.session_state.llm_analysis = None
                st.session_state.llm_dataset = None
        with c2:
            if st.session_state.llm_analysis:
                if st.button("保存本次校对经验", key="llm_save_exp"):
                    _save_experience(llm_file_name, st.session_state.llm_analysis)
                    st.success("校对经验已保存，下次加载同类文件时可复用")


def _load_and_update_preview(tmp_path, data_id, analysis, progress=None):
    from lmbagent.data.llm_loader import llm_smart_load

    analysis_lines = []
    analysis_lines.append(f"**识别格式:** `{analysis.get('format', 'unknown')}`")
    analysis_lines.append(f"**分隔符:** `{analysis.get('separator', ',')}`")
    analysis_lines.append(f"**跳过行数:** {analysis.get('skip_rows', 0)}")
    analysis_lines.append(f"**编码:** `{analysis.get('encoding', 'utf-8')}`")

    if analysis.get("column_map"):
        analysis_lines.append("**列映射:**")
        for std, actual in analysis["column_map"].items():
            if actual:
                analysis_lines.append(f"  `{std}` ← `{actual}`")

    if analysis.get("reasoning"):
        analysis_lines.append(f"\n*AI说明: {analysis['reasoning']}*")

    with st.spinner("正在加载数据..."):
        try:
            ds = llm_smart_load(
                str(tmp_path),
                data_id=data_id or None,
                analysis=analysis,
            )
            if progress:
                progress.progress(100, text="加载完成!")
                progress.empty()
            if ds and ds.num_data_points > 0:
                st.session_state.llm_dataset = ds
                col_info = ", ".join(ds.raw_data.columns.tolist())
                st.session_state.llm_messages.append({
                    "role": "assistant",
                    "content": (
                        f"AI 加载完成!\n\n"
                        + "\n".join(analysis_lines) + "\n\n"
                        f"**数据:** {ds.num_data_points} 行, {ds.num_cycles} 循环\n"
                        f"**列:** {col_info}\n\n"
                        "请检查上方数据预览，如需调整请在下方输入校正指令。"
                    ),
                })
            else:
                st.session_state.llm_messages.append({
                    "role": "assistant",
                    "content": (
                        "AI 分析完成但加载结果为空。\n\n"
                        + "\n".join(analysis_lines) + "\n\n"
                        "请在下方描述文件格式，我会帮你校正加载。"
                    ),
                })
        except Exception as e:
            if progress:
                progress.empty()
            st.session_state.llm_messages.append({
                "role": "assistant",
                "content": (
                    f"AI 加载失败: {e}\n\n"
                    + "\n".join(analysis_lines) + "\n\n"
                    "请在下方描述文件格式，我会帮你校正加载。"
                ),
            })


def _apply_experience_and_load(tmp_path, data_id, analysis):
    st.session_state.llm_messages.append({
        "role": "assistant",
        "content": f"已应用历史校对经验 (format={analysis.get('format')})，正在重新加载...",
    })
    _load_and_update_preview(tmp_path, data_id, analysis)
    st.rerun()


def _render_data_preview(ds_preview):
    with st.container():
        st.markdown("#### 数据预览")
        c1, c2, c3 = st.columns(3)
        c1.metric("数据行数", f"{ds_preview.num_data_points:,}")
        c2.metric("循环数", ds_preview.num_cycles)
        c3.metric("列数", len(ds_preview.raw_data.columns))
        st.markdown(f"**列名:** `{list(ds_preview.raw_data.columns)}`")

        df_prev = ds_preview.raw_data
        cc_col = df_prev.get("charge_capacity")
        dc_col = df_prev.get("discharge_capacity")
        cc_sum = cc_col.sum() if cc_col is not None else 0
        dc_sum = dc_col.sum() if dc_col is not None else 0
        if cc_sum > 0 or dc_sum > 0:
            st.success(f"充电容量总和: {cc_sum:.4f} Ah | 放电容量总和: {dc_sum:.4f} Ah")
        else:
            st.warning("充电容量和放电容量总和均为 0 — 数据可能缺少容量列或需要按电流拆分")

        active_rows = df_prev[df_prev["current"] != 0] if "current" in df_prev.columns else df_prev.head(20)
        st.dataframe(
            active_rows.head(20) if len(active_rows) > 0 else df_prev.head(20),
            hide_index=True,
            use_container_width=True,
        )
        if ds_preview.num_cycles > 0 and dc_col is not None:
            cyc = df_prev.groupby("cycle_index")["discharge_capacity"].max()
            cyc_nonzero = cyc[cyc > 0]
            if len(cyc_nonzero) > 0:
                st.markdown(f"**放电容量范围:** {cyc_nonzero.min():.4f} ~ {cyc_nonzero.max():.4f} Ah (非零值)")


# --- Experience persistence ---

def _save_experience(file_name: str, analysis: dict):
    _EXPERIENCE_DIR.mkdir(parents=True, exist_ok=True)
    exp = {
        "name": file_name,
        "analysis": analysis,
    }
    out = _EXPERIENCE_DIR / f"{Path(file_name).stem}.json"
    out.write_text(json.dumps(exp, ensure_ascii=False, indent=2), encoding="utf-8")


def _list_experiences() -> list[dict]:
    if not _EXPERIENCE_DIR.exists():
        return []
    results = []
    for p in sorted(_EXPERIENCE_DIR.glob("*.json")):
        try:
            d = json.loads(p.read_text(encoding="utf-8"))
            results.append(d)
        except Exception:
            pass
    return results


def _render_batch_upload(store: DataStore):
    st.subheader("批量上传")
    st.markdown("支持两种方式批量上传：从服务器目录路径批量加载，或逐个文件上传。")

    use_ai = st.checkbox("对失败文件启用 AI 智能加载", value=True, key="batch_use_ai")

    split_mode = st.selectbox(
        "多循环文件处理",
        ["不拆分 (整个文件作为一个数据集)", "按循环拆分 (每个循环独立)", "按 N 个循环分组"],
        key="batch_split_mode",
    )
    split_n = 1
    if split_mode == "按循环拆分 (每个循环独立)":
        split_n = 1
    elif split_mode == "按 N 个循环分组":
        split_n = st.number_input("每组循环数", min_value=2, value=50, key="batch_split_n")
    else:
        split_n = 0

    source = st.radio("数据来源", ["服务器目录路径", "逐个文件上传"], key="batch_source", horizontal=True)

    if source == "服务器目录路径":
        batch_dir = st.text_input(
            "输入服务器目录路径",
            placeholder="/path/to/experiment/data",
            key="batch_dir",
        )

        if batch_dir and Path(batch_dir).is_dir():
            extensions = {".csv", ".xlsx", ".npy", ".txt", ".xls", ".tsv"}
            all_files = sorted(
                p for p in Path(batch_dir).rglob("*")
                if p.suffix.lower() in extensions and not p.name.startswith("~") and p.stat().st_size > 0
            )

            if all_files:
                st.write(f"发现 {len(all_files)} 个数据文件:")
                formats = {}
                for f in all_files:
                    fmt = "未知"
                    try:
                        fmt = detect_format(str(f))
                    except Exception:
                        pass
                    formats[fmt] = formats.get(fmt, 0) + 1
                st.markdown("**格式统计:** " + ", ".join(f"`{k}`: {v}" for k, v in formats.items()))
                with st.expander("文件列表"):
                    for f in all_files:
                        st.markdown(f"- `{f.name}` ({f.stat().st_size / 1024:.1f} KB)")

                if st.button("批量导入", type="primary", key="batch_dir_import"):
                    _do_batch_import(store, all_files, use_ai, split_n)
            else:
                st.warning("目录下未发现数据文件")
        elif batch_dir:
            st.error(f"目录不存在: {batch_dir}")

    else:
        files = st.file_uploader(
            "逐个添加数据文件（每个文件独立上传，无总大小限制）",
            type=["csv", "xlsx", "npy", "txt"],
            accept_multiple_files=True,
            key="multi_upload",
        )

        if files:
            st.write(f"已选择 {len(files)} 个文件:")
            for f in files:
                fmt = "未知"
                try:
                    with tempfile.NamedTemporaryFile(suffix=f.name, delete=False) as tmp:
                        tmp.write(f.getvalue())
                        fmt = detect_format(tmp.name)
                except Exception:
                    pass
                st.markdown(f"  - `{f.name}` ({fmt}, {f.size / 1024:.1f} KB)")

            if st.button("批量导入", type="primary", key="batch_file_import"):
                tmp_files = []
                for f in files:
                    tmpdir = tempfile.mkdtemp()
                    tmp_path = Path(tmpdir) / f.name
                    tmp_path.write_bytes(f.getvalue())
                    tmp_files.append(tmp_path)
                _do_batch_import(store, tmp_files, use_ai, split_n)


def _do_batch_import(store, file_paths, use_ai, split_n):
    from lmbagent.data.llm_loader import llm_analyze_file, llm_smart_load

    progress = st.progress(0)
    status_container = st.container()
    success_count = 0
    ai_count = 0
    fail_count = 0
    split_count = 0
    results = []

    for i, fpath in enumerate(file_paths):
        fpath = Path(fpath) if not isinstance(fpath, Path) else fpath
        status = "⏳"
        loaded = False

        try:
            ds = load_auto(str(fpath))
            if ds.num_data_points > 0:
                if split_n > 0 and ds.num_cycles > split_n:
                    sub_datasets = split_by_cycles(ds, max_cycles_per_dataset=split_n)
                    for sub_ds in sub_datasets:
                        sub_ds = add_cycle_summary(sub_ds)
                        store.put(sub_ds)
                    split_count += len(sub_datasets) - 1
                else:
                    ds = add_cycle_summary(ds)
                    store.put(ds)
                success_count += 1
                status = "✅"
                loaded = True
        except Exception:
            pass

        if not loaded and use_ai:
            try:
                analysis = llm_analyze_file(str(fpath))
                ds = llm_smart_load(str(fpath), analysis=analysis)
                if ds and ds.num_data_points > 0:
                    if split_n > 0 and ds.num_cycles > split_n:
                        sub_datasets = split_by_cycles(ds, max_cycles_per_dataset=split_n)
                        for sub_ds in sub_datasets:
                            sub_ds = add_cycle_summary(sub_ds)
                            store.put(sub_ds)
                        split_count += len(sub_datasets) - 1
                    else:
                        ds = add_cycle_summary(ds)
                        store.put(ds)
                    ai_count += 1
                    status = "🤖"
                    loaded = True
            except Exception:
                pass

        if not loaded:
            fail_count += 1
            status = "❌"

        results.append({"file": fpath.name, "status": status, "loaded": loaded})

        progress.progress((i + 1) / len(file_paths))

    with status_container:
        cols = st.columns(5)
        cols[0].metric("总计", len(file_paths))
        cols[1].metric("✅ 标准加载", success_count)
        cols[2].metric("🤖 AI 加载", ai_count)
        cols[3].metric("❌ 失败", fail_count)
        cols[4].metric("🔀 拆分数据集", split_count)

        with st.expander("详细结果"):
            for r in results:
                st.markdown(f"{r['status']} `{r['file']}`")

    st.metric("数据库中实验总数", len(store.list_ids()))


def _render_directory_scan(store: DataStore):
    st.subheader("目录扫描与 AI 智能探索")
    st.markdown("""
    扫描文件服务器上的目录，自动发现并批量导入实验数据。
    支持 AI 智能分析目录结构和命名模式。
    """)

    scan_dir = st.text_input(
        "输入目录路径",
        placeholder="/path/to/experiment/data",
        key="scan_dir",
    )

    use_ai = st.checkbox("对加载失败文件启用 AI 智能加载", value=True, key="scan_use_ai")

    split_mode = st.selectbox(
        "多循环文件处理",
        ["不拆分 (整个文件作为一个数据集)", "按循环拆分 (每个循环独立)", "按 N 个循环分组"],
        key="scan_split_mode",
    )
    split_n = 1
    if split_mode == "按循环拆分 (每个循环独立)":
        split_n = 1
    elif split_mode == "按 N 个循环分组":
        split_n = st.number_input("每组循环数", min_value=2, value=50, key="scan_split_n")
    else:
        split_n = 0

    if scan_dir and Path(scan_dir).is_dir():
        col_scan, col_ai = st.columns(2)
        with col_scan:
            scan_btn = st.button("扫描目录", key="scan_btn")
        with col_ai:
            ai_scan_btn = st.button("AI 智能探索目录", type="primary", key="ai_scan_btn")

        if scan_btn:
            with st.spinner("正在扫描..."):
                candidates = scan_directory(scan_dir)
            if not candidates:
                st.warning("未发现数据文件。")
            else:
                st.session_state.scan_candidates = candidates
                _display_scan_results(candidates)

        if ai_scan_btn:
            with st.spinner("AI 正在分析目录结构..."):
                try:
                    from lmbagent.data.catalog import smart_scan_directory
                    result = smart_scan_directory(scan_dir)
                except Exception as e:
                    st.error(f"AI 分析失败: {e}")
                    result = None

            if result:
                st.session_state.ai_scan_result = result

                if result.get("summary"):
                    st.info(f"AI 分析: {result['summary']}")

                if result.get("patterns"):
                    st.subheader("识别的模式")
                    for p in result["patterns"]:
                        st.markdown(f"- {p}")

                if result.get("groups"):
                    st.subheader("文件分组")
                    for group_name, files in result["groups"].items():
                        with st.expander(f"{group_name} ({len(files)} 文件)"):
                            for f in files:
                                st.markdown(f"  - `{f}`")

                if result.get("candidates"):
                    st.session_state.scan_candidates = result["candidates"]
                    _display_scan_results(result["candidates"])

        if "scan_candidates" in st.session_state and st.session_state.scan_candidates:
            candidates = st.session_state.scan_candidates
            if st.button(f"导入全部 {len(candidates)} 个文件", key="import_all_btn"):
                from lmbagent.data.llm_loader import llm_analyze_file, llm_smart_load

                progress = st.progress(0)
                success_count = 0
                ai_count = 0
                fail_count = 0
                split_count = 0

                with st.spinner("批量导入中..."):
                    for i, c in enumerate(candidates):
                        loaded = False
                        try:
                            ds = load_auto(str(c.path))
                            if c.design and ds.experiment_design is None:
                                ds.experiment_design = c.design
                            if ds.num_data_points > 0:
                                if split_n > 0 and ds.num_cycles > split_n:
                                    sub_datasets = split_by_cycles(ds, max_cycles_per_dataset=split_n)
                                    for sub_ds in sub_datasets:
                                        if c.design and sub_ds.experiment_design is None:
                                            sub_ds.experiment_design = c.design
                                        sub_ds = add_cycle_summary(sub_ds)
                                        store.put(sub_ds)
                                    split_count += len(sub_datasets) - 1
                                else:
                                    ds = add_cycle_summary(ds)
                                    store.put(ds)
                                c.imported = True
                                c.data_id = ds.data_id
                                success_count += 1
                                loaded = True
                        except Exception:
                            pass

                        if not loaded and use_ai:
                            try:
                                analysis = llm_analyze_file(str(c.path))
                                ds = llm_smart_load(str(c.path), analysis=analysis)
                                if ds and ds.num_data_points > 0:
                                    if c.design and ds.experiment_design is None:
                                        ds.experiment_design = c.design
                                    if split_n > 0 and ds.num_cycles > split_n:
                                        sub_datasets = split_by_cycles(ds, max_cycles_per_dataset=split_n)
                                        for sub_ds in sub_datasets:
                                            if c.design and sub_ds.experiment_design is None:
                                                sub_ds.experiment_design = c.design
                                            sub_ds = add_cycle_summary(sub_ds)
                                            store.put(sub_ds)
                                        split_count += len(sub_datasets) - 1
                                    else:
                                        ds = add_cycle_summary(ds)
                                        store.put(ds)
                                    c.imported = True
                                    c.data_id = ds.data_id
                                    ai_count += 1
                                    loaded = True
                            except Exception:
                                pass

                        if not loaded:
                            fail_count += 1

                        progress.progress((i + 1) / len(candidates))

                c1, c2, c3, c4 = st.columns(4)
                c1.metric("✅ 标准加载", success_count)
                c2.metric("🤖 AI 加载", ai_count)
                c3.metric("❌ 失败", fail_count)
                c4.metric("🔀 拆分数据集", split_count)
                st.metric("数据库中实验总数", len(store.list_ids()))

    elif scan_dir:
        st.error(f"目录不存在: {scan_dir}")


def _display_scan_results(candidates):
    st.write(f"发现 {len(candidates)} 个数据文件:")
    formats = {}
    for c in candidates:
        design_info = " ✓ 有设计元数据" if c.design_path else ""
        st.markdown(f"  - `{c.path.name}` ({c.format}){design_info}")
        formats[c.format] = formats.get(c.format, 0) + 1

    if formats:
        st.markdown("**格式统计:** " + ", ".join(f"{k}: {v}" for k, v in formats.items()))


def _render_smart_explore(store: DataStore):
    st.subheader("智能探索加载")
    st.markdown("""
    上传格式未知或混乱的数据文件，系统会自动探索文件结构并尝试结构化。
    适用于非标准格式、混合编码、不规则表头等场景。
    """)

    split_mode = st.selectbox(
        "多循环文件处理",
        ["不拆分 (整个文件作为一个数据集)", "按循环拆分 (每个循环独立)", "按 N 个循环分组"],
        key="explore_split_mode",
    )
    split_n = 1
    if split_mode == "按循环拆分 (每个循环独立)":
        split_n = 1
    elif split_mode == "按 N 个循环分组":
        split_n = st.number_input("每组循环数", min_value=2, value=50, key="explore_split_n")
    else:
        split_n = 0

    explore_file = st.file_uploader(
        "上传数据文件",
        type=["csv", "txt", "tsv", "xlsx", "xls"],
        key="explore_upload",
    )

    explore_server_path = st.text_input(
        "或输入服务器文件路径",
        placeholder="/path/to/your/data.csv",
        key="explore_server_path",
    )

    tmp_path = None
    file_name = None

    if explore_server_path and Path(explore_server_path).exists():
        tmp_path = Path(explore_server_path)
        file_name = tmp_path.name
        st.write(f"**文件:** {file_name} ({tmp_path.stat().st_size / 1024:.1f} KB)")
    elif explore_file:
        tmp_dir = Path(tempfile.gettempdir()) / "lmbagent_explore"
        tmp_dir.mkdir(parents=True, exist_ok=True)
        tmp_path = tmp_dir / explore_file.name
        tmp_path.write_bytes(explore_file.getvalue())
        file_name = explore_file.name
        st.write(f"**文件:** {file_name}")

    if tmp_path:
        with st.expander("文件探索结果", expanded=True):
            _explore_file(tmp_path)

        data_id = st.text_input("数据集 ID", value="", key="explore_id",
                                placeholder=file_name.split(".")[0] if file_name else "")

        col1, col2 = st.columns(2)
        with col1:
            if st.button("尝试加载", key="explore_load"):
                try:
                    ds = load_auto(str(tmp_path), data_id=data_id or None)
                    _store_dataset(store, ds, split_n)
                except Exception as e:
                    st.error(f"自动加载失败: {e}")
        with col2:
            if st.button("AI 智能加载", key="explore_ai_load"):
                try:
                    from lmbagent.data.llm_loader import llm_analyze_file, llm_smart_load
                    with st.spinner("AI 正在分析..."):
                        analysis = llm_analyze_file(str(tmp_path))
                        ds = llm_smart_load(str(tmp_path), data_id=data_id or None, analysis=analysis)
                    if ds and ds.num_data_points > 0:
                        _store_dataset(store, ds, split_n)
                        if analysis.get("reasoning"):
                            st.info(f"AI: {analysis['reasoning']}")
                    else:
                        st.error("AI 加载失败")
                except Exception as e:
                    st.error(f"AI 加载失败: {e}")


def _store_dataset(store, ds, split_n):
    if ds.num_data_points == 0:
        st.error("加载数据为空")
        return

    if split_n > 0 and ds.num_cycles > split_n:
        sub_datasets = split_by_cycles(ds, max_cycles_per_dataset=split_n)
        for sub_ds in sub_datasets:
            sub_ds = add_cycle_summary(sub_ds)
            store.put(sub_ds)
        st.success(
            f"加载成功并拆分为 {len(sub_datasets)} 个数据集! "
            f"(原: {ds.num_data_points} points, {ds.num_cycles} cycles)"
        )
    else:
        ds = add_cycle_summary(ds)
        store.put(ds)
        st.success(f"加载成功! ID: `{ds.data_id}`, {ds.num_data_points} points, {ds.num_cycles} cycles")
    st.session_state.active_dataset_id = ds.data_id





def _show_file_preview(path: Path):
    suffix = path.suffix.lower()
    if suffix in (".csv", ".txt", ".tsv"):
        for enc in ["utf-8", "gbk", "latin-1"]:
            try:
                with open(path, "r", encoding=enc, errors="replace") as f:
                    lines = [f.readline().rstrip() for _ in range(10)]
                for i, line in enumerate(lines):
                    st.code(f"L{i+1}: {line[:200]}", language=None)
                break
            except Exception:
                continue
    elif suffix in (".xlsx", ".xls"):
        try:
            import openpyxl
            wb = openpyxl.load_workbook(str(path), data_only=True)
            st.markdown(f"**Sheets:** {wb.sheetnames}")
            for sn in wb.sheetnames[:3]:
                ws = wb[sn]
                rows = list(ws.iter_rows(values_only=True, max_row=6))
                if rows:
                    st.markdown(f"**Sheet '{sn}'**:")
                    for i, row in enumerate(rows):
                        st.code(
                            f"L{i+1}: {[str(v) if v is not None else '' for v in row][:10]}",
                            language=None,
                        )
            wb.close()
        except Exception as e:
            st.error(f"Excel 读取失败: {e}")
    elif suffix == ".npy":
        st.info("二进制 numpy 文件，无法文本预览")


def _explore_file(path: Path):
    suffix = path.suffix.lower()
    st.markdown(f"**扩展名:** `{suffix}`")
    st.markdown(f"**大小:** {path.stat().st_size / 1024:.1f} KB")

    if suffix in (".csv", ".txt", ".tsv"):
        try:
            with open(path, "rb") as f:
                raw = f.read(4096)

            encodings = ["utf-8", "gbk", "gb2312", "latin-1", "utf-16"]
            detected_enc = None
            for enc in encodings:
                try:
                    raw.decode(enc)
                    detected_enc = enc
                    break
                except (UnicodeDecodeError, ValueError):
                    continue
            st.markdown(f"**编码:** {detected_enc or 'unknown'}")

            if detected_enc:
                with open(path, "r", encoding=detected_enc, errors="replace") as f:
                    lines = [f.readline().rstrip() for _ in range(10)]
                st.markdown("**前10行:**")
                for i, line in enumerate(lines):
                    st.code(f"L{i+1}: {line[:200]}", language=None)

            seps = [",", "\t", ";", "|"]
            best_sep = ","
            best_cols = 0
            for sep in seps:
                try:
                    df = pd.read_csv(path, sep=sep, nrows=5, encoding=detected_enc or "utf-8",
                                     on_bad_lines="skip")
                    if len(df.columns) > best_cols:
                        best_cols = len(df.columns)
                        best_sep = sep
                except Exception:
                    pass

            st.markdown(f"**分隔符:** `{repr(best_sep)}` ({best_cols} 列)")

            try:
                df = pd.read_csv(path, sep=best_sep, nrows=5, encoding=detected_enc or "utf-8",
                                 on_bad_lines="skip")
                st.markdown(f"**列名:** {list(df.columns)}")
                st.dataframe(df.head(), width="stretch", hide_index=True)

                battery_keywords = ["voltage", "current", "capacity", "cycle", "time",
                                    "voltage(v)", "current(a)", "capacity(ah)",
                                    "电压", "电流", "容量", "循环", "时间"]
                found = [c for c in df.columns if any(kw in str(c).lower() for kw in battery_keywords)]
                if found:
                    st.success(f"检测到电池数据列: {found}")
                else:
                    st.warning("未检测到标准电池数据列名")
            except Exception as e:
                st.warning(f"解析失败: {e}")

        except Exception as e:
            st.error(f"无法读取文件: {e}")

    elif suffix in (".xlsx", ".xls"):
        try:
            import zipfile
            with zipfile.ZipFile(str(path)) as z:
                import re
                wb_xml = z.read("xl/workbook.xml").decode("utf-8")
                sheets = re.findall(r'name="([^"]+)"', wb_xml)
            st.markdown(f"**Sheet列表:** {sheets}")
            import openpyxl
            wb = openpyxl.load_workbook(str(path), data_only=True)
            for sn in sheets[:3]:
                ws = wb[sn]
                rows = list(ws.iter_rows(values_only=True, max_row=6))
                if rows:
                    st.markdown(f"**Sheet '{sn}' 前6行:**")
                    for i, row in enumerate(rows):
                        st.code(f"L{i+1}: {[str(v) if v is not None else '' for v in row][:10]}",
                                language=None)
            wb.close()
        except Exception as e:
            st.error(f"Excel读取失败: {e}")
